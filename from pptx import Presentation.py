from pptx import Presentation
from pptx.util import Inches
import os
import pandas as pd


def position_image(pptx_file, image_path, top_in, left_in):
    """
    Adds an image to a PowerPoint slide at the specified position in inches.

    Args:
        pptx_file (str): Path to the PowerPoint file.
        image_path (str): Path to the image file.
        top_in (float): Top position of the image in inches.
        left_in (float): Left position of the image in inches.
    """
    try:
        prs = Presentation(pptx_file)
        slide = prs.slides[0]  # Assuming you want to add to the first slide

        # Calculate the top and left position in EMUs (English Metric Units)
        top = Inches(top_in)
        left = Inches(left_in)

        # Add the image to the slide
        slide.shapes.add_picture(image_path, left, top)

        prs.save(pptx_file)
        print(
            f"Image added successfully to {pptx_file} at top={top_in}in, left={left_in}in"
        )

    except Exception as e:
        print(f"An error occurred: {e}")


if __name__ == "__main__":
    # Load the CSV file
    csv_file = "URLfanfav.csv"  # Replace with your CSV file name
    pptx_folder = "PPTX_Files"  # Folder containing PPTX files
    image_folder = "QR_Poster_Images"  # Folder containing image files

    # Read the CSV file into a DataFrame
    df = pd.read_csv(csv_file)

    # Iterate through each row in the DataFrame
    for index, row in df.iterrows():
        submission_id = row["Submission ID"]
        poster_id = row["POSTER_ID"]

        # Construct file paths
        pptx_file = os.path.join(pptx_folder, f"{submission_id}.pptx")
        image_path = os.path.join(image_folder, f"{poster_id}.png")

        # Check if both files exist
        if os.path.exists(pptx_file) and os.path.exists(image_path):
            # Set the position for the image
            top_in = 0.63
            left_in = 32.02

            # Call the function to add the image to the PPTX file
            position_image(pptx_file, image_path, top_in, left_in)
        else:
            print(
                f"Skipping: Missing file for Submission ID {submission_id} or POSTER_ID {poster_id}"
            )
